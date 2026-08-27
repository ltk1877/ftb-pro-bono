import csv
import re
from datetime import datetime
from pathlib import Path

import openpyxl

SRC = "dashboard-inputs.xlsx"
OUT_DIR = Path("csv")

# Sheet3 has no header row; it's a ZIP -> County -> Census Tract -> City lookup table.
SHEET_NAME_OVERRIDES = {
    "Sheet3": "zip_county_tract_city_lookup",
}

# Sheets with no header row in the source workbook.
HEADERLESS_SHEETS = {
    "Sheet3": ["zip", "county", "census_tract_geoid", "city"],
}


def slugify(name: str) -> str:
    name = name.strip()
    name = re.sub(r"[^\w\-]+", "_", name)
    name = re.sub(r"_+", "_", name).strip("_")
    return name.lower()


def slugify_header(value) -> str:
    text = "" if value is None else str(value).strip()

    # Collapse midnight timestamps (used as monthly column headers) to
    # month_yyyy_mm. Prefixed so it's a valid unquoted SQL identifier
    # (Snowflake/most SQL dialects don't allow identifiers to start with a digit).
    try:
        dt = datetime.strptime(text, "%Y-%m-%d %H:%M:%S")
        if dt.hour == dt.minute == dt.second == 0:
            return f"month_{dt.year:04d}_{dt.month:02d}"
    except ValueError:
        pass

    text = text.replace("%", " pct ")
    text = text.replace("#", " num ")
    text = text.replace("&", " and ")
    text = re.sub(r"[^\w]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text.lower()


def clean_headers(raw_headers: list) -> list:
    headers = [slugify_header(h) for h in raw_headers]
    seen: dict[str, int] = {}
    deduped = []
    for h in headers:
        h = h or "col"
        if h in seen:
            seen[h] += 1
            h = f"{h}_{seen[h]}"
        else:
            seen[h] = 0
        deduped.append(h)
    return deduped


def trim_trailing_blank_columns(rows: list) -> list:
    if not rows:
        return rows
    last_col = -1
    for row in rows:
        for i, v in enumerate(row):
            if v is not None and str(v).strip() != "":
                last_col = max(last_col, i)
    return [row[: last_col + 1] for row in rows]


def main() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        display_name = SHEET_NAME_OVERRIDES.get(sheet_name, sheet_name)
        out_path = OUT_DIR / f"{slugify(display_name)}.csv"

        rows = list(ws.iter_rows(values_only=True))
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        rows = trim_trailing_blank_columns(rows)

        if sheet_name in HEADERLESS_SHEETS:
            header = HEADERLESS_SHEETS[sheet_name]
            data_rows = rows
        else:
            header = clean_headers(rows[0]) if rows else []
            data_rows = rows[1:]

        with out_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(header)
            for row in data_rows:
                writer.writerow(["" if v is None else v for v in row])
        print(f"{sheet_name!r} -> {out_path} ({len(data_rows)} rows)")


if __name__ == "__main__":
    main()
