# -*- coding: utf-8 -*-
"""Input validation for the new-year scoring tool (scripts/score_new_year.py).

Fails fast with a specific, readable message instead of letting a bad input file surface
as a cryptic KeyError deep inside feature engineering.
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.features.build_features import build_transition_table


class ValidationError(Exception):
    """Raised when an input file doesn't have what the pipeline needs, with a specific reason."""


REQUIRED_ZCTA_COLUMNS = [
    "County FIPS", "ZCTA", "Geography", "County, State", "State", "Year",
    "Food Bank 1 ID", "Food Bank 1", "Total Population (5 Year ACS)",
    "Overall Food Insecurity Rate", "# of Food Insecure Persons Overall",
    "Unemployment Rate (1 Yr BLS)", "Poverty Rate (5 Yr ACS)", "Percent Black (5 Yr ACS)",
    "Percent Hispanic (any race) (5 Year ACS)", "Median Income (5 Yr ACS)",
    "Homeownership Rate (5 Yr ACS)", "Disability Rate (5 Yr ACS)",
]

REQUIRED_COUNTY_COLUMNS = [
    "FIPS", "Year", "Total Population (5 Year ACS)", "Total Child Population (5 Year ACS)",
    "Percent White, non-Hispanic (5 Year ACS)", "Cost Per Meal", "Weighted Index",
    "SNAP Threshold", "Rural-Urban Continuum Code (2023)",
]

REQUIRED_ALICE_COLUMNS = [
    "Year", "GEO id2", "ALICE Households", "Poverty Households", "Households",
    "ALICE Threshold - HH under 65", "ALICE Threshold - HH 65 years and over",
]


def require_file(path: Path, description: str) -> Path:
    path = Path(path)
    if not path.exists():
        raise ValidationError(f"{description} not found: {path}")
    return path


def validate_workbook_sheets(path: Path, required_sheets: list) -> None:
    wb = load_workbook(path, read_only=True)
    missing = [s for s in required_sheets if s not in wb.sheetnames]
    if missing:
        raise ValidationError(
            f"{path.name} is missing sheet(s) {missing}. "
            f"Sheets found in this workbook: {wb.sheetnames}"
        )


def validate_columns(df: pd.DataFrame, required_cols: list, context: str) -> None:
    missing = [c for c in required_cols if c not in df.columns]
    if missing:
        raise ValidationError(
            f"{context} is missing required column(s): {missing}. "
            f"Columns found: {list(df.columns)}"
        )


def resolve_source_year(df: pd.DataFrame, year_col: str = "Year", override: int = None) -> int:
    """Pick the single year this data represents, or fail with the ambiguity spelled out."""
    years = sorted(pd.to_numeric(df[year_col], errors="coerce").dropna().unique().tolist())
    years = [int(y) for y in years]

    if override is not None:
        if override not in years:
            raise ValidationError(
                f"--source-year {override} not found in this data. Years present: {years}"
            )
        return override

    if len(years) == 0:
        raise ValidationError(f"No usable values found in the '{year_col}' column.")
    if len(years) > 1:
        raise ValidationError(
            f"This file contains multiple years {years} -- pass --source-year to pick one "
            f"(the tool expects a single-year snapshot, e.g. one MMG release)."
        )
    return years[0]


def validate_minimum_coverage(panel: pd.DataFrame, source_year: int, state: str, driver_cols: list, minimum_rows: int = 10) -> int:
    """Confirm there's enough complete data at the source year (for the target state, if given) to forecast from."""
    at_year = panel.loc[panel["year"] == source_year].copy()
    if state:
        at_year = at_year.loc[at_year["State"].astype(str).str.upper() == state.upper()]
    complete = at_year.dropna(subset=["food_insecurity_rate"] + driver_cols)

    if len(complete) < minimum_rows:
        scope = f"for state '{state}' " if state else ""
        raise ValidationError(
            f"Only {len(complete)} ZIP rows {scope}have complete data at year {source_year} "
            f"(need at least {minimum_rows}). Check that the source file actually covers "
            f"{scope}and year {source_year}."
        )
    return len(complete)


def validate_minimum_training_rows(panel: pd.DataFrame, driver_lag: int, rate_lag: int, driver_cols: list, minimum_rows: int = 10) -> int:
    """Confirm there are enough complete (target-year, lagged-year) transitions to fit a model.

    Unlike `validate_minimum_coverage` (one source year's snapshot, for scoring), training needs
    paired observations `minimum_rows` rows deep -- built the same way `evaluate_lag_spec` builds
    them -- so a bad or too-thin input surfaces here instead of as a raw ValueError/LinAlgError.
    """
    table = build_transition_table(panel, driver_lag, rate_lag, driver_cols)
    complete = table.dropna(subset=["food_insecurity_rate", "lag_food_insecurity_rate"] + driver_cols)

    if len(complete) < minimum_rows:
        raise ValidationError(
            f"Only {len(complete)} complete (year, year-{rate_lag}) transitions found across the "
            f"loaded data (need at least {minimum_rows}). This usually means the loaded MMG "
            f"workbook(s) don't span enough years -- a {rate_lag}-year lag needs rows with both a "
            f"target year and an observation {rate_lag} years earlier for the same ZIP."
        )
    return len(complete)
